from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "Spej_OS_Unified_Operations_Handoff.docx"
SCREENSHOTS = ROOT / "output" / "playwright"

NAVY = "1A2332"
DEEP_NAVY = "111A27"
CARD = "1D2838"
TEAL = "4ECDC4"
TEAL_DARK = "0E7A72"
INK = "1A2332"
MUTED = "5F6B7C"
FAINT = "8792A2"
LINE = "D6DEE8"
LIGHT = "EDF1F5"
TEAL_LIGHT = "DDF7F4"
WHITE = "F8F9FA"
GOLD = "D7A63C"

FONT = "Calibri"

REQUIRED_SCREENSHOTS = {
    "my_work": SCREENSHOTS / "spej-my-work-dark.png",
    "gtm": SCREENSHOTS / "spej-gtm-dark.png",
    "projects": SCREENSHOTS / "spej-projects-dark.png",
    "sosa": SCREENSHOTS / "spej-sosa-meeting-dark.png",
    "access": SCREENSHOTS / "spej-access-dark.png",
    "integrations": SCREENSHOTS / "spej-integrations-dark.png",
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, *, size: float = 11, color: str = INK, bold: bool = False, italic: bool = False, font: str = FONT) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], *, indent: int = 120, borders: bool = True) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is not None:
        tbl_pr.remove(tbl_borders)
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single" if borders else "nil")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LINE)
        tbl_borders.append(node)
    tbl_pr.append(tbl_borders)


def set_paragraph(paragraph, *, before: float = 0, after: float = 6, line: float = 1.25, align=None, keep: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_with_next = keep
    if align is not None:
        paragraph.alignment = align


def add_text(doc_or_cell, text: str, *, size: float = 11, color: str = INK, bold: bool = False, italic: bool = False, before: float = 0, after: float = 6, line: float = 1.25, align=None, keep: bool = False):
    paragraph = doc_or_cell.add_paragraph() if hasattr(doc_or_cell, "add_paragraph") else None
    if paragraph is None:
        raise RuntimeError("Unsupported paragraph container")
    set_paragraph(paragraph, before=before, after=after, line=line, align=align, keep=keep)
    set_run(paragraph.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return paragraph


def replace_cell_text(cell, text: str, *, size: float = 10, color: str = INK, bold: bool = False, after: float = 2, line: float = 1.15, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_paragraph(paragraph, after=after, line=line, align=align)
    set_run(paragraph.add_run(text), size=size, color=color, bold=bold)


def add_heading(doc, text: str, level: int = 1) -> None:
    style = doc.styles[f"Heading {level}"]
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    if level == 1:
        set_run(run, size=16, color=TEAL_DARK, bold=True)
    elif level == 2:
        set_run(run, size=13, color=TEAL_DARK, bold=True)
    else:
        set_run(run, size=12, color=NAVY, bold=True)


def create_bullet_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•")
    level.append(text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL_DARK)
    r_pr.append(color)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, text: str, num_id: int, *, bold_lead: str | None = None, size: float = 10.3) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, after=4, line=1.25)
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(num)
    p_pr.append(num_pr)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    if bold_lead and text.startswith(bold_lead):
        set_run(paragraph.add_run(bold_lead), size=size, color=INK, bold=True)
        set_run(paragraph.add_run(text[len(bold_lead):]), size=size, color=INK)
    else:
        set_run(paragraph.add_run(text), size=size, color=INK)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL_DARK)
    run_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(underline)
    run.append(run_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    set_run(run, size=8.5, color=FAINT)


def configure_document(doc: Document) -> int:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, before, after, color in ((1, 16, 18, 10, TEAL_DARK), (2, 13, 14, 7, TEAL_DARK), (3, 12, 10, 5, NAVY)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [4680, 4680], indent=0, borders=False)
    replace_cell_text(table.cell(0, 0), "SPEJ OS", size=8.5, color=TEAL_DARK, bold=True, after=0)
    replace_cell_text(table.cell(0, 1), "UNIFIED OPERATIONS PREVIEW", size=8.5, color=FAINT, bold=True, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    header.paragraphs[0]._element.getparent().remove(header.paragraphs[0]._element)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(footer_table, [4680, 4680], indent=0, borders=False)
    replace_cell_text(footer_table.cell(0, 0), "Internal working brief · pilot", size=8.5, color=FAINT, after=0)
    page_paragraph = footer_table.cell(0, 1).paragraphs[0]
    page_paragraph.clear()
    set_paragraph(page_paragraph, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_run(page_paragraph.add_run("Page "), size=8.5, color=FAINT)
    add_page_number(page_paragraph)
    footer.paragraphs[0]._element.getparent().remove(footer.paragraphs[0]._element)
    return create_bullet_numbering(doc)


def add_brand_cover(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=0, borders=False)
    cell = table.cell(0, 0)
    set_cell_fill(cell, NAVY)
    set_cell_margins(cell, top=200, start=260, bottom=200, end=260)
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    set_paragraph(paragraph, after=8, line=1.0)
    paragraph.add_run().add_picture(str(ROOT / "public" / "spej-logo.png"), width=Inches(0.9))
    title = cell.add_paragraph()
    set_paragraph(title, after=4, line=1.0)
    set_run(title.add_run("Spej OS"), size=28, color=WHITE, bold=True)
    subtitle = cell.add_paragraph()
    set_paragraph(subtitle, after=0, line=1.1)
    set_run(subtitle.add_run("Unified company operations · GTM, CRM, Projects, and SOSA"), size=12.5, color=TEAL, bold=True)
    add_text(doc, "Working pilot and IT handoff · September 3, 2026", size=9.5, color=MUTED, after=10)


def add_status_strip(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=120, borders=True)
    set_cell_fill(table.cell(0, 0), TEAL_LIGHT)
    replace_cell_text(table.cell(0, 0), text, size=9.5, color=TEAL_DARK, bold=True, after=0)
    add_text(doc, "", size=2, after=2)


def add_screenshot(doc: Document, key: str, caption: str, *, width: float = 6.15) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, before=3, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
    paragraph.add_run().add_picture(str(REQUIRED_SCREENSHOTS[key]), width=Inches(width))
    caption_paragraph = doc.add_paragraph()
    set_paragraph(caption_paragraph, before=0, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(caption_paragraph.add_run(caption), size=8.5, color=MUTED, italic=True)


def add_workspace_table(doc: Document) -> None:
    rows = [
        ("My Work", "Role-relevant tasks, deadlines, approvals, updates, and exceptions."),
        ("CRM", "One company-wide record for accounts, people, client history, opportunities, partners, and activity."),
        ("GTM", "Focused sales and marketing workspace: pipeline, partners, content, campaigns, work, metrics, and intelligence."),
        ("Projects", "Company execution for client work, AI Office, Plooms, events, media, product, partner, and internal projects."),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    set_table_geometry(table, [1800, 7560])
    for index, label in enumerate(("Workspace", "Primary use")):
        set_cell_fill(table.cell(0, index), NAVY)
        replace_cell_text(table.cell(0, index), label, size=9.5, color=WHITE, bold=True, after=0)
    for row_index, (label, detail) in enumerate(rows, start=1):
        set_cell_fill(table.cell(row_index, 0), LIGHT)
        replace_cell_text(table.cell(row_index, 0), label, size=9.5, color=INK, bold=True, after=0)
        replace_cell_text(table.cell(row_index, 1), detail, size=9.2, color=INK, after=0)


def add_ai_table(doc: Document) -> None:
    rows = [
        ("Identity, permissions, and field access", "Summaries, extraction, and classification"),
        ("Required fields, state transitions, and duplicate checks", "Meeting-to-record proposals and draft follow-ups"),
        ("Scoring math, metrics, sync cursors, retries, and idempotency", "Priority explanations and missing-context questions"),
        ("Approved commits, audit logs, and service health", "Content drafts and account preparation"),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    set_table_geometry(table, [4680, 4680])
    for index, label in enumerate(("Deterministic system", "LLM-assisted work")):
        set_cell_fill(table.cell(0, index), NAVY if index == 0 else TEAL_DARK)
        replace_cell_text(table.cell(0, index), label, size=9.5, color=WHITE, bold=True, after=0)
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            if row_index % 2 == 0:
                set_cell_fill(table.cell(row_index, col_index), LIGHT)
            replace_cell_text(table.cell(row_index, col_index), value, size=9, color=INK, after=0)


def add_three_column_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    set_table_geometry(table, [2600, 3380, 3380])
    for index, label in enumerate(("Foundation", "Retain", "Add or map")):
        set_cell_fill(table.cell(0, index), NAVY)
        replace_cell_text(table.cell(0, index), label, size=9.2, color=WHITE, bold=True, after=0)
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            if col_index == 0:
                set_cell_fill(table.cell(row_index, col_index), LIGHT)
            replace_cell_text(table.cell(row_index, col_index), value, size=8.7, color=INK, bold=col_index == 0, after=0)


def add_architecture_stack(doc: Document) -> None:
    rows = [
        ("1 · WHERE PEOPLE WORK", "SOSA web · Microsoft Teams · mobile · dashboard"),
        ("2 · GOVERNANCE", "Entra identity · roles and grants · tool allowlists · approval · audit"),
        ("3 · SYSTEMS OF RECORD", "Spej CRM and Projects · tickets and QA · documents and knowledge · Microsoft 365"),
    ]
    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [2450, 6910], indent=120, borders=True)
    fills = [TEAL_DARK, NAVY, CARD]
    for index, (label, detail) in enumerate(rows):
        set_cell_fill(table.cell(index, 0), fills[index])
        set_cell_fill(table.cell(index, 1), LIGHT if index < 2 else "E4EAF0")
        replace_cell_text(table.cell(index, 0), label, size=9.5, color=WHITE, bold=True, after=0)
        replace_cell_text(table.cell(index, 1), detail, size=10, color=INK, bold=index == 0, after=0)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> Path:
    missing = [str(path) for path in REQUIRED_SCREENSHOTS.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing final screenshots:\n" + "\n".join(missing))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    bullet_num_id = configure_document(doc)

    # Page 1 — summary overview and key bullets.
    add_brand_cover(doc)
    add_status_strip(doc, "DEMO READY · Local interface, workflows, test data, and contracts. NOT CONNECTED · Production Spej OS data, SOSA, Microsoft 365, identity, or company security services.")
    add_heading(doc, "Summary overview", 1)
    add_bullet(doc, "One shared record target. The pilot uses one local model; production maps CRM, GTM, Projects, tasks, content, and metrics to existing Spej canonical IDs.", bullet_num_id, bold_lead="One shared record target.")
    add_bullet(doc, "Role-based My Work. Each person starts with relevant work; separate permissions control data and actions.", bullet_num_id, bold_lead="Role-based My Work.")
    add_bullet(doc, "SOSA-first operation. The pilot supports local questions, pasted meetings, reviewable proposals, and local approved actions; production tools remain an IT integration.", bullet_num_id, bold_lead="SOSA-first operation.")
    add_bullet(doc, "Original ownership boundaries preserved in the integration design. In production, Microsoft 365, documents and knowledge, tickets and QA, connectors, audit, feature controls, model settings, and usage services remain authoritative.", bullet_num_id, bold_lead="Original ownership boundaries preserved in the integration design.")
    add_bullet(doc, "GTM is a focused view, not a second system. It adds sales, partner, content, metrics, and intelligence workflows over shared records.", bullet_num_id, bold_lead="GTM is a focused view, not a second system.")
    add_screenshot(doc, "my_work", "My Work: role-relevant tasks, SOSA entry, and linked operating views.", width=5.55)

    # Page 2 — operating model.
    page_break(doc)
    add_heading(doc, "How the system works", 1)
    add_text(doc, "Spej OS remains the governed company system. SOSA is the primary interaction layer; the dashboard is the review, exception, and administration layer.", size=10.5, after=9)
    add_architecture_stack(doc)
    add_heading(doc, "Operating rules", 2)
    add_bullet(doc, "Store each account, person, opportunity, project, task, ticket, file reference, and decision once.", bullet_num_id)
    add_bullet(doc, "Filter every search, SOSA response, action, and linked file by the signed-in user’s effective access.", bullet_num_id)
    add_bullet(doc, "Use human approval for sensitive writes; allow low-risk, reversible actions only under explicit policy.", bullet_num_id)
    add_bullet(doc, "Show source coverage and connector health so SOSA never assumes an incomplete knowledge source is complete.", bullet_num_id)
    add_text(doc, "Demo path: My Work → ask SOSA or process a transcript → review linked CRM, GTM, and Project records → show People & access and Integrations as the production handoff.", size=9.4, color=TEAL_DARK, bold=True, before=4, after=6)
    add_heading(doc, "Industry analogy", 2)
    p = add_text(doc, "The architecture is similar in concept—not affiliation—to Salesforce and Anthropic’s Claudeforce: an AI interface over governed enterprise data, workflows, business logic, and actions. Spej’s version uses SOSA over Spej OS and Microsoft 365.", size=9.8, color=MUTED, after=3)
    source = doc.add_paragraph()
    set_paragraph(source, before=4, after=4, line=1.0)
    set_run(source.add_run("Source: "), size=8.5, color=FAINT, bold=True)
    add_hyperlink(source, "Salesforce and Anthropic announce Claudeforce, August 26, 2026", "https://www.salesforce.com/news/press-releases/2026/08/26/salesforce-and-anthropic-announce-claudeforce/")

    # Page 3 — workspace logic and GTM.
    page_break(doc)
    add_heading(doc, "The workspace logic", 1)
    add_workspace_table(doc)
    add_text(doc, "The existing canonical Spej CRM service remains the company source of truth. GTM and Projects become purpose-built views over shared records. My Work is the person-level action view.", size=9.8, color=MUTED, before=5, after=4)
    add_screenshot(doc, "gtm", "GTM: pipeline, partner work, content, activity, metrics, and intelligence in one focused workspace.", width=5.55)

    # Page 4 — Projects and retained project controls.
    page_break(doc)
    add_heading(doc, "Projects and delivery", 1)
    add_bullet(doc, "Supports client work, AI Office, Plooms, events, marketing and media, partner enablement, product, and internal initiatives.", bullet_num_id)
    add_bullet(doc, "Discovery, Design, and Delivery is one optional playbook—not the definition of every project.", bullet_num_id)
    add_bullet(doc, "Won opportunities are checked for missing projects; project work supports owners, priorities, subtasks, and separate deadlines.", bullet_num_id)
    add_bullet(doc, "Reference slots are provided for tickets, QA, governed files, transcripts, decisions, approvals, time entries, and enterprise records; authorized records appear after their existing service adapters are connected.", bullet_num_id)
    add_screenshot(doc, "projects", "Projects: portfolio health, execution work, sale-to-project checks, and connected project records.", width=5.75)

    # Page 5 — SOSA and deterministic/LLM split.
    page_break(doc)
    add_heading(doc, "SOSA and meeting-driven updates", 1)
    add_text(doc, "The pilot demonstrates pasted-transcript intake. After production connectors are added, Teams transcripts, voice notes, and chat requests can enter the same reviewable workflow: capture → extract → link evidence → validate → approve → save → audit.", size=10, after=5)
    add_screenshot(doc, "sosa", "Local design preview: transcript context becomes proposed CRM, GTM, Project, and task updates. No live Teams or SOSA connector is used in this pilot.", width=5.05)
    add_ai_table(doc)
    add_text(doc, "The LLM interprets and drafts. Deterministic services control identity, permissions, validation, state, commits, synchronization, and audit.", size=9.2, color=MUTED, before=4, after=2)

    # Page 6 — access and administration.
    page_break(doc)
    add_heading(doc, "Access and administration", 1)
    add_bullet(doc, "Default views come from job role; access is separately granted by tenant, portal, action, data scope, record, field, feature, and time window.", bullet_num_id)
    add_bullet(doc, "Preserve existing global_admin, super_admin, user, and service_department role IDs, then approve a crosswalk to the narrower preview policies.", bullet_num_id)
    add_bullet(doc, "Confidential records require exact grants. Temporary GTM, contractor, API, MCP, and service-account access must expire or be revocable.", bullet_num_id)
    add_bullet(doc, "SOSA, search, knowledge retrieval, exports, and Microsoft file links must apply the same permission check before retrieval and before action.", bullet_num_id)
    add_screenshot(doc, "access", "Local access-design preview: role, portal, scope, actions, exact records, expiry, and audit. Production policy must be persisted and enforced by existing Spej services.", width=5.65)

    # Page 7 — retained foundations and integrations.
    page_break(doc)
    add_heading(doc, "Current Spej OS foundations this design reuses", 1)
    add_three_column_table(doc, [
        ("Identity and admin", "Tenant isolation; existing global_admin/super_admin/user/service_department IDs; access duration; feature controls.", "Approve a role crosswalk, then bind the shell to signed-in Spej identity and server-side authorization."),
        ("Microsoft 365", "Teams meetings/chats/files; Outlook mail/calendar/contacts; SharePoint and OneDrive.", "Use Microsoft Graph events and sync checkpoints, stable external IDs, and governed deep links."),
        ("Knowledge", "Document processing, vector indexing, transcript metadata, knowledge graph, permission-aware retrieval.", "Add artifact/version contracts, citations, ACL intersection, deletion/re-index, and source health."),
        ("Operations", "CRM; project grid, Kanban, Gantt, milestones and time; tickets and QA; enterprise scorecards, goals, issues, meetings; API/MCP, webhooks and playbooks.", "Show authorized references in My Work, CRM, GTM, Projects, and SOSA."),
        ("Control plane", "Audit logs, job health, retries, feature flags, model controls, usage and cost reporting.", "Expose status/deep links; do not rebuild the control plane in the GTM shell."),
        ("Brand and generation", "White-label settings, Spej branding, mobile packaging, and presentation generation.", "Retain existing controls and call approved services through governed SOSA tools."),
    ])
    add_text(doc, "Also retained: Google Workspace, file-share, REST and Azure SQL connector options. The detailed compatibility matrix is the no-regression checklist; no current screen or service is removed until it is mapped, tested, and accepted.", size=8.9, color=MUTED, before=4, after=3)
    add_screenshot(doc, "integrations", "Integration contract: the new interface reuses existing Spej OS services instead of creating a second source of truth.", width=5.3)

    # Page 8 — implementation and definition of ready.
    page_break(doc)
    add_heading(doc, "Implementation path", 1)
    add_bullet(doc, "1. Create a Spej-owned repository. Keep it private or internally approved unless leadership completes a public-release review.", bullet_num_id, bold_lead="1. Create a Spej-owned repository.")
    add_bullet(doc, "2. Confirm canonical service ownership and IDs for CRM, Projects, tickets, quality, decisions, files, knowledge, features, and usage.", bullet_num_id, bold_lead="2. Confirm canonical service ownership and IDs")
    add_bullet(doc, "3. Connect the existing Spej identity and authorization services, using Entra federation where approved; verify denied users cannot access portals, records, fields, tools, exports, or knowledge results.", bullet_num_id, bold_lead="3. Connect the existing Spej identity and authorization services,")
    add_bullet(doc, "4. Add Microsoft Graph and existing Spej API/MCP adapters with verified events, sync checkpoints, duplicate-safe retries, a failed-event queue, and audit.", bullet_num_id, bold_lead="4. Add Microsoft Graph and existing Spej API/MCP adapters")
    add_bullet(doc, "5. Register the versioned tools with the existing SOSA. Keep SOSA’s model and runtime under the approved Spej architecture; do not use a personal ChatGPT subscription as production runtime.", bullet_num_id, bold_lead="5. Register the versioned tools with the existing SOSA.")
    add_bullet(doc, "6. Run a role-based pilot, migrate only validated data, compare results with the existing Spej OS, and retain rollback.", bullet_num_id, bold_lead="6. Run a role-based pilot,")

    add_heading(doc, "Definition of ready for company use", 2)
    readiness = [
        "Production starts only when database, identity, audit, queue, Microsoft, SOSA, telemetry, and existing Spej service registry checks pass.",
        "CRM, project, ticket, quality, document, transcript, decision, time, enterprise, feature, and usage parity mappings are signed off.",
        "Negative permission tests cover UI, API, SOSA, MCP, search, knowledge retrieval, files, exports, and confidential records.",
        "Meeting automation preserves transcript/artifact provenance on every derived record and never invents commitments.",
        "Connector coverage, job failures, usage, latency, and cost are visible to authorized administrators.",
        "Backup, restore, retention, deletion, incident response, and rollback have been tested before broader rollout.",
    ]
    for item in readiness:
        add_bullet(doc, item, bullet_num_id, size=9.8)

    add_text(doc, "Production guard: this repository refuses to start in production mode until IT replaces local demo stores and reference adapters with authenticated, durable company services.", size=9.2, color=TEAL_DARK, bold=True, before=3, after=3)
    add_text(doc, "Security note: the security material is a design-time threat model—not a penetration test, compliance certification, or production approval. Production requires server-side access checks, permission-filtered model context, reviewed writes, audit, secret management, backup/restore, and connector kill switches. SOSA must not receive database or Microsoft Graph credentials.", size=8.8, color=MUTED, after=4)

    add_heading(doc, "Repository handoff index", 2)
    add_text(doc, "Start with README.md and docs/GITHUB_HANDOFF.md. Then review docs/EXISTING_SPEJ_OS_COMPATIBILITY.md, DATA_OWNERSHIP_MATRIX.md, SOSA_TOOL_CONTRACT.md, MICROSOFT_INTEGRATION_CONTRACT.md, KNOWLEDGE_AND_DOCUMENT_INTEGRATION_CONTRACT.md, ACCESS_CONTROL_AND_TASK_VISIBILITY.md, SECURITY_THREAT_MODEL.md, DEPLOYMENT_RUNBOOK.md, and MIGRATION_AND_ROLLBACK.md. Preserve the repository LICENSE and required copyright notice.", size=8.9, color=MUTED, after=0)

    doc.core_properties.title = "Spej OS — Unified Company Operations Preview"
    doc.core_properties.subject = "Concise product overview, demo guide, and IT handoff"
    doc.core_properties.author = "Spej"
    doc.core_properties.keywords = "Spej OS, SOSA, CRM, GTM, Projects, Microsoft 365, handoff"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
